import pytest
from docx import Document

from docusense.ingestion.chunking import chunk_pages
from docusense.ingestion.loaders import EmptyDocument, Page, UnsupportedFileType, load_document
from docusense.text import split_sentences, tokenize
from tests.conftest import SAMPLES


def test_tokenize_strips_punctuation_and_stopwords():
    assert tokenize("What is DocuSense?") == ["docusense"]
    assert tokenize("Self-driving cars, V2G!") == ["self-driving", "cars", "v2g"]


def test_split_sentences_handles_bullets_without_periods():
    text = "Key Features:\n- Secure offline processing\n- BM25 retrieval\nIt works well. Really."
    assert split_sentences(text) == [
        "Key Features:",
        "Secure offline processing",
        "BM25 retrieval",
        "It works well.",
        "Really.",
    ]


def test_load_pdf_has_page_numbers():
    pages = load_document(SAMPLES / "cloud_mobility_proposal.pdf")
    assert len(pages) > 3
    assert pages[0].number == 1
    assert "Automotive" in pages[0].text


def test_load_docx(tmp_path):
    path = tmp_path / "notes.docx"
    doc = Document()
    doc.add_paragraph("First paragraph about retrieval.")
    doc.add_paragraph("Second paragraph about generation.")
    doc.save(path)
    pages = load_document(path)
    assert "retrieval" in pages[0].text and "generation" in pages[0].text


def test_unsupported_and_empty_files(tmp_path):
    bad = tmp_path / "x.csv"
    bad.write_text("a,b")
    with pytest.raises(UnsupportedFileType):
        load_document(bad)
    empty = tmp_path / "empty.txt"
    empty.write_text("   \n  ")
    with pytest.raises(EmptyDocument):
        load_document(empty)


def test_chunking_respects_size_overlap_and_pages():
    sentences = [f"Sentence number {i} talks about topic {i}." for i in range(60)]
    pages = [Page(text=" ".join(sentences[:30]), number=1), Page(text=" ".join(sentences[30:]), number=2)]
    chunks = chunk_pages(pages, "doc1", "f.pdf", chunk_size=50, chunk_overlap=12)

    assert all(len(c.text.split()) <= 50 for c in chunks)
    assert {c.page for c in chunks} == {1, 2}
    assert [c.position for c in chunks] == list(range(len(chunks)))
    # consecutive chunks on the same page share their boundary sentence
    same_page = [(a, b) for a, b in zip(chunks, chunks[1:], strict=False) if a.page == b.page]
    assert same_page and all(a.text.splitlines()[-1] in b.text for a, b in same_page)


def test_chunking_splits_huge_sentence():
    pages = [Page(text=" ".join(["word"] * 500))]
    chunks = chunk_pages(pages, "d", "f.txt", chunk_size=100, chunk_overlap=0)
    assert len(chunks) == 5


def test_chunking_rejects_bad_params():
    with pytest.raises(ValueError):
        chunk_pages([Page(text="hi")], "d", "f", chunk_size=10, chunk_overlap=10)


def test_split_sentences_handles_missing_space_after_period():
    assert split_sentences("Follows GDPR and CCPA standards.Data is encrypted. Use Node.js daily.") == [
        "Follows GDPR and CCPA standards.",
        "Data is encrypted.",
        "Use Node.js daily.",
    ]
