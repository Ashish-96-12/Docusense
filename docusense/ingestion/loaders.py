"""Turn PDF, DOCX, TXT and Markdown files into a list of pages."""
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

from docusense.text import normalize_whitespace

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFileType(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


@dataclass
class Page:
    text: str
    number: Optional[int] = None  # 1-based, only for PDFs


def load_document(path: Path) -> List[Page]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == ".pdf":
        pages = _load_pdf(path)
    elif ext == ".docx":
        pages = _load_docx(path)
    else:
        pages = [Page(text=path.read_text(encoding="utf-8", errors="replace"))]

    pages = [Page(text=normalize_whitespace(p.text), number=p.number) for p in pages]
    pages = [p for p in pages if p.text]
    if not pages:
        raise EmptyDocument(
            f"No text found in {path.name}. If it's a scanned PDF it needs OCR first."
        )
    return pages


def _load_pdf(path: Path) -> List[Page]:
    import logging

    from pypdf import PdfReader

    # pypdf logs harmless warnings for slightly malformed PDFs
    logging.getLogger("pypdf").setLevel(logging.ERROR)

    reader = PdfReader(str(path))
    return [Page(text=page.extract_text() or "", number=i + 1) for i, page in enumerate(reader.pages)]


def _load_docx(path: Path) -> List[Page]:
    from docx import Document

    doc = Document(str(path))
    blocks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Page(text="\n\n".join(b for b in blocks if b.strip()))]
